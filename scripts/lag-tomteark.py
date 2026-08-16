#!/usr/bin/env python3
"""Genererer docs/tomtedetaljer-sandmoen.xlsx fra data/tomter.json + data/config.json.

Arbeidsboka er et gjennomgangs- og utfyllingsark for Frode: alle detaljer per tomt
samlet ett sted, med gule celler der nettsiden mangler innhold i dag. Svarene hans
skrives tilbake til JSON-filene av utvikler.

Kjør:  python3 scripts/lag-tomteark.py
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "tomtedetaljer-sandmoen.xlsx"

tomter = json.loads((ROOT / "data" / "tomter.json").read_text(encoding="utf-8"))
config = json.loads((ROOT / "data" / "config.json").read_text(encoding="utf-8"))
fokus = config.get("fokus", {})

GRONN = "26412F"
SAND = "F4F0E6"
LINJE = "E2DCCB"
GUL = "FFF3CD"
GRA = "8A917F"
PLASSHOLDER = "Mer informasjon kommer."

FONT = "Arial"
h_font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
h_fill = PatternFill("solid", fgColor=GRONN)
b_font = Font(name=FONT, size=10)
b_fill = PatternFill("solid", fgColor=GUL)
sand_fill = PatternFill("solid", fgColor=SAND)
tittel_font = Font(name=FONT, size=14, bold=True, color=GRONN)
dempet = Font(name=FONT, size=10, italic=True, color=GRA)
kant = Border(*[Side(style="thin", color=LINJE)] * 4)
topp = Alignment(vertical="top", wrap_text=True)
topp_midt = Alignment(vertical="top", horizontal="center", wrap_text=True)


def sett_bredder(ws, bredder):
    for kol, bredde in bredder.items():
        ws.column_dimensions[kol].width = bredde


def skriv_hode(ws, rad, overskrifter, hoyde=34):
    for i, tekst in enumerate(overskrifter, start=1):
        c = ws.cell(row=rad, column=i, value=tekst)
        c.font, c.fill, c.alignment, c.border = h_font, h_fill, topp, kant
    ws.row_dimensions[rad].height = hoyde


def mangler(vare):
    """Felt som står tomme eller er plassholdertekst på nettsiden i dag."""
    ut = []
    if not vare.get("terreng") or vare["terreng"] == PLASSHOLDER:
        ut.append("terreng")
    return ut


def mangler_formel(rad):
    """Liste over manglende felt, satt sammen med & og IF så den virker i enhver Excel-versjon."""
    deler = (
        f'IF(OR($E{rad}="",$E{rad}="{PLASSHOLDER}"),"Beskrivelse · ","")'
        f'&IF($G{rad}=0,"Bilder · ","")'
    )
    return f'=IF(LEN({deler})=0,"Komplett",LEFT({deler},LEN({deler})-3))'


# ---------------------------------------------------------------- Tomter
wb = Workbook()
ws = wb.active
ws.title = "Tomter"

KOLONNER = [
    ("Tomt", 7),
    ("BFR i\nreguleringsplan", 14),
    ("Status", 12),
    ("Areal\n(dekar)", 9),
    ("Beskrivelse av tomta\n(ingress øverst på tomtesiden)", 52),
    ("Plassering i feltet\n(tekst over reguleringskartet)", 42),
    ("Bilder\n(antall)", 9),
    ("Video", 8),
    ("Tekst under videoen", 26),
    ("Kart\nX %", 7),
    ("Kart\nY %", 7),
    ("Mangler på nettsiden i dag", 26),
    ("FRODE: tilføyelser og rettelser", 46),
    ("FRODE:\ngjennomgått?", 14),
]

ws["A1"] = "Tomtedetaljer · Sandmoen hyttetomter"
ws["A1"].font = tittel_font
ws["A2"] = (
    "Alt som ligger på sandmoen.com per tomt i dag. Gule celler mangler innhold. "
    "Skriv rett i de to siste kolonnene, eller rett direkte i cellene – begge deler funker. "
    "Se arket «Veiledning» for hva hvert felt styrer."
)
ws["A2"].font = dempet
ws.merge_cells("A2:N2")
ws.row_dimensions[2].height = 30
ws["A2"].alignment = topp

skriv_hode(ws, 4, [k for k, _ in KOLONNER])
sett_bredder(ws, {chr(64 + i): b for i, (_, b) in enumerate(KOLONNER, start=1)})

rad = 5
for t in tomter:
    m = mangler(t)
    verdier = [
        t["nr"],
        t.get("bfr", ""),
        t["status"],
        t["areal"],
        t.get("terreng", ""),
        t.get("plassering", ""),
        f'=COUNTIFS(Bilder!$A$5:$A$400,$A{rad},Bilder!$B$5:$B$400,"Bilde")',  # kolonne G
        "Ja" if t.get("video") else "Nei",
        t.get("videoTekst", ""),
        t.get("x", ""),
        t.get("y", ""),
        mangler_formel(rad),
        "",
        "",
    ]
    for i, v in enumerate(verdier, start=1):
        c = ws.cell(row=rad, column=i, value=v)
        c.font, c.alignment, c.border = b_font, topp, kant
        if i in (1, 3, 4, 7, 8, 10, 11, 14):
            c.alignment = topp_midt
    ws.cell(row=rad, column=4).number_format = "0.0"
    if "terreng" in m:
        ws.cell(row=rad, column=5).fill = b_fill
    ws.cell(row=rad, column=13).fill = b_fill
    ws.cell(row=rad, column=14).fill = b_fill
    ws.row_dimensions[rad].height = 118
    rad += 1

siste = rad - 1
ws.cell(row=rad + 1, column=1, value="Eksempel på utfylling – ikke en ekte tomt, slett gjerne raden:").font = dempet
ws.merge_cells(start_row=rad + 1, start_column=1, end_row=rad + 1, end_column=8)
eks_rad = rad + 2
eksempel = [
    99, "BFR99", "Ledig", 1.1,
    "Slak, tørr furumo med fjell i dagen mot vest. Bilvei fram til P3, derfra ca 80 meter på opparbeidet sti.",
    "Ligger mellom BFR 17 og BFR 18, merket bBF i reguleringskartet.",
    0, "Nei", "", 50, 50, "", "Tar nye bilder til høsten.", "Ja",
]
for i, v in enumerate(eksempel, start=1):
    c = ws.cell(row=eks_rad, column=i, value=v)
    c.font, c.alignment, c.border = dempet, topp, kant
    c.fill = sand_fill
    if i in (1, 3, 4, 7, 8, 10, 11, 14):
        c.alignment = topp_midt
ws.cell(row=eks_rad, column=4).number_format = "0.0"
ws.cell(row=eks_rad, column=7).value = 0
ws.row_dimensions[eks_rad].height = 76

ws.cell(row=4, column=12).comment = Comment(
    "Regnes ut automatisk fra kolonnene til venstre. «Komplett» betyr at "
    "beskrivelse og minst ett bilde er på plass.",
    "Sandmoen",
)

dv_status = DataValidation(type="list", formula1='"Ledig,Reservert,Festet"', allow_blank=True)
dv_janei = DataValidation(type="list", formula1='"Ja,Nei"', allow_blank=True)
ws.add_data_validation(dv_status)
ws.add_data_validation(dv_janei)
dv_status.add(f"C5:C{siste}")
dv_janei.add(f"N5:N{siste}")

ws.freeze_panes = "B5"
ws.auto_filter.ref = f"A4:N{siste}"
ws.sheet_view.zoomScale = 90

# ---------------------------------------------------------------- Bilder
wsb = wb.create_sheet("Bilder")
wsb["A1"] = "Bilder og video per tomt"
wsb["A1"].font = tittel_font
wsb["A2"] = (
    "Rekkefølgen bestemmer visningen: bilde 1 er hovedbildet på tomtesiden og på kortet i oversikten, "
    "bilde 2 og 3 vises som miniatyrer, resten ligger bak «+ N bilder» i bildevisningen."
)
wsb["A2"].font = dempet
wsb.merge_cells("A2:H2")
wsb.row_dimensions[2].height = 28
wsb["A2"].alignment = topp

skriv_hode(wsb, 4, [
    "Tomt", "Type", "Nr", "Filnavn", "Vises som", "Bildeutsnitt",
    "FRODE: beholde?", "FRODE: kommentar eller bildetekst",
])
sett_bredder(wsb, {"A": 7, "B": 13, "C": 6, "D": 30, "E": 30, "F": 16, "G": 15, "H": 48})

rb = 5
dv_beholde = DataValidation(type="list", formula1='"Ja,Nei,Byttes"', allow_blank=True)
wsb.add_data_validation(dv_beholde)
for t in tomter:
    bilder = t.get("bilder", []) or []
    for i, fil in enumerate(bilder):
        if i == 0:
            rolle = "Hovedbilde (tomtesiden + kort i oversikten)"
        elif i < 3:
            rolle = "Miniatyr ved siden av hovedbildet"
        else:
            rolle = "Kun i bildevisningen («+ N bilder»)"
        for kol, v in enumerate([t["nr"], "Bilde", i + 1, fil, rolle, fokus.get(fil, "standard"), "", ""], start=1):
            c = wsb.cell(row=rb, column=kol, value=v)
            c.font, c.alignment, c.border = b_font, topp, kant
            if kol in (1, 2, 3, 7):
                c.alignment = topp_midt
        wsb.cell(row=rb, column=7).fill = b_fill
        wsb.cell(row=rb, column=8).fill = b_fill
        rb += 1
    if t.get("video"):
        for typ, fil, rolle in (
            ("Video", t["video"], "Knappen «Se video fra tomta»"),
            ("Video-plakat", t.get("videoPoster", ""), "Stillbildet før videoen starter"),
        ):
            for kol, v in enumerate([t["nr"], typ, "", fil, rolle, "", "", ""], start=1):
                c = wsb.cell(row=rb, column=kol, value=v)
                c.font, c.alignment, c.border = b_font, topp, kant
                c.fill = sand_fill
                if kol in (1, 2, 3, 7):
                    c.alignment = topp_midt
            wsb.cell(row=rb, column=7).fill = b_fill
            wsb.cell(row=rb, column=8).fill = b_fill
            rb += 1

dv_beholde.add(f"G5:G{rb - 1}")
wsb.freeze_panes = "A5"
wsb.auto_filter.ref = f"A4:H{rb - 1}"

# ---------------------------------------------------------------- Felles
wsf = wb.create_sheet("Felles for alle tomter")
wsf["A1"] = "Priser, vilkår og tekster som gjelder alle tomtene"
wsf["A1"].font = tittel_font
wsf["A2"] = "Endres disse, endres de samtidig på alle sidene."
wsf["A2"].font = dempet

skriv_hode(wsf, 4, ["Hva", "Verdi i dag", "Hvor det vises", "FRODE: endring"])
sett_bredder(wsf, {"A": 32, "B": 44, "C": 42, "D": 42})

vilkaar = config["vilkaar"]
kontakt = config["kontakt"]
sted = config["sted"]

felles = [
    ("Engangsbeløp ved feste", vilkaar["engangsbelop"], "Forside, oversikt, hver tomteside, kontakt"),
    ("Årlig festeavgift", vilkaar["festeavgift"], "Forside, oversikt, hver tomteside, kontakt"),
    ("Kommunale kostnader", vilkaar["kommunaleKostnader"], "Oversikten («Hyttetomter til feste»)"),
    ("Merknad om vann, vei og strøm", config["merknad"], "Oversikt, hver tomteside, kontakt"),
]
for i, punkt in enumerate(config["folgerMed"], start=1):
    felles.append((f"«Dette følger med» – punkt {i}", punkt, "Hver tomteside"))
felles += [
    ("Tekst under hero-videoen", config.get("heroVideoTekst", ""), "Forsiden"),
    ("Stedsnavn", sted["navn"], "Topplinje og bunntekst på alle sider"),
    ("Undertittel", sted["undertittel"], "Topplinje og bunntekst"),
    ("Feltnavn", sted["felt"], "Forsiden og hver tomteside"),
    ("Kontaktperson", kontakt["navn"], "Bunntekst og kontaktsiden"),
    ("Telefon", kontakt["telefon"], "Bunntekst og kontaktsiden"),
    ("E-post", kontakt["epost"], "Bunntekst og kontaktsiden"),
    ("Lenke til reguleringsplanen", config["lenker"]["reguleringsplan"], "Oversikten og hver tomteside"),
    ("Lenke til tomtefesteloven", config["lenker"]["tomtefesteloven"], "«Slik fester du tomt» på oversikten"),
    ("Lenke om nasjonalparken", config["lenker"]["nasjonalpark"], "Oversikten og hver tomteside"),
]

rf = 5
for hva, verdi, hvor in felles:
    for kol, v in enumerate([hva, verdi, hvor, ""], start=1):
        c = wsf.cell(row=rf, column=kol, value=v)
        c.font, c.alignment, c.border = b_font, topp, kant
    if isinstance(verdi, (int, float)):
        wsf.cell(row=rf, column=2).number_format = '#,##0 "kr"'
    wsf.cell(row=rf, column=4).fill = b_fill
    wsf.row_dimensions[rf].height = 30
    rf += 1

wsf.freeze_panes = "A5"

# ---------------------------------------------------------------- Veiledning
wsv = wb.create_sheet("Veiledning")
wsv["A1"] = "Slik bruker du arket"
wsv["A1"].font = tittel_font
sett_bredder(wsv, {"A": 30, "B": 62, "C": 44})

intro = [
    "Arket viser alt nettsiden vet om hver tomt i dag – ingenting er skjult i koden utenom dette.",
    "Gule celler er innhold som mangler eller er en plassholder. Det er de som haster.",
    "Du kan enten rette rett i cellene, eller skrive i FRODE-kolonnene. Jan Erik legger det inn på nettsiden.",
    "Raden merket «Eksempel» nederst på arket «Tomter» viser formatet – den er ikke en ekte tomt.",
    "Arket er et øyeblikksbilde. Etter at endringene er lagt inn på nettsiden lages arket på nytt.",
]
rv = 3
for linje in intro:
    c = wsv.cell(row=rv, column=1, value="•")
    c.font, c.alignment = b_font, topp_midt
    wsv.column_dimensions["A"].width = 30
    c2 = wsv.cell(row=rv, column=2, value=linje)
    c2.font, c2.alignment = b_font, topp
    wsv.merge_cells(start_row=rv, start_column=2, end_row=rv, end_column=3)
    rv += 1

rv += 1
wsv.cell(row=rv, column=1, value="Fargekoder").font = Font(name=FONT, size=11, bold=True, color=GRONN)
rv += 1
for farge, tekst in ((GUL, "Mangler innhold, eller er din å fylle ut"), (SAND, "Eksempel eller filnavn – kun til orientering")):
    c = wsv.cell(row=rv, column=1, value="")
    c.fill = PatternFill("solid", fgColor=farge)
    c.border = kant
    c2 = wsv.cell(row=rv, column=2, value=tekst)
    c2.font, c2.alignment = b_font, topp
    rv += 1

rv += 1
skriv_hode(wsv, rv, ["Felt", "Hva det er", "Hvor på nettsiden"])
rv += 1

forklaringer = [
    ("Tomt", "Tomtenummeret. Brukes i adressen til tomtesiden (sandmoen.com/tomt/7/).", "Overalt"),
    ("BFR i reguleringsplan", "Betegnelsen i reguleringsplanen. Følger tomtenummeret på alle tomtene i dag.", "Avsnittet «Plassering i feltet»"),
    ("Status", "Ledig, Reservert eller Festet. «Festet» vises som «Bortfestet» på nettsiden.", "Merke på kort, tomteside og i filteret"),
    ("Areal", "Størrelse i dekar, én desimal (1,0 da).", "Kort, faktaboks og prispanel"),
    ("Beskrivelse av tomta", "Ingressen øverst på tomtesiden, og teksten på kortet i oversikten. To–fire setninger om terreng, adkomst, strøm og hva som gjør tomta spesiell.", "Tomtesiden og oversiktskortet"),
    ("Plassering i feltet", "Setningen over reguleringskartet, som forklarer hvor tomta ligger i forhold til naboene. Står den tom, skrives «I reguleringsplanen er dette BFR N» automatisk.", "Tomtesiden"),
    ("Bilder", "Antall bilder på tomta. Detaljene ligger på arket «Bilder».", "Galleri på tomtesiden"),
    ("Video", "Om tomta har video. Alle fem har sommervideo fra 2026.", "Knappen «Se video fra tomta»"),
    ("Tekst under videoen", "Valgfri bildetekst i videovinduet, f.eks. «Tomt 14 · F4 grense nord · Høst».", "Videovinduet"),
    ("Kart X og Y", "Hvor tomtenålen står i kartvisningen i oversikten, målt i prosent av bredde og høyde. Bare relevant hvis nålen står feil.", "Kartvisningen i oversikten"),
    ("Mangler på nettsiden i dag", "Regnes ut automatisk. «Komplett» betyr at beskrivelse og bilder er på plass.", "Vises ikke på nettsiden"),
]
for felt, hva, hvor in forklaringer:
    for kol, v in enumerate([felt, hva, hvor], start=1):
        c = wsv.cell(row=rv, column=kol, value=v)
        c.font, c.alignment, c.border = b_font, topp, kant
    wsv.row_dimensions[rv].height = 42
    rv += 1

rv += 1
wsv.cell(row=rv, column=1, value="Ikke med i arket").font = Font(name=FONT, size=11, bold=True, color=GRONN)
rv += 1
for tekst in (
    "Feltet har 15 tomter totalt. Bare de fem som er til feste nå (7, 8, 14, 17 og 18) ligger på nettsiden og i dette arket. Skal flere legges ut, sier du fra, så føyer vi dem til.",
    "Bilder legges ikke inn i arket – de bearbeides fra originalfilene. Skriv i kommentarfeltet hvilke som skal byttes, så tar vi det derfra.",
):
    c = wsv.cell(row=rv, column=1, value=tekst)
    c.font, c.alignment = b_font, topp
    wsv.merge_cells(start_row=rv, start_column=1, end_row=rv, end_column=3)
    wsv.row_dimensions[rv].height = 30
    rv += 1

for sheet in wb.worksheets:
    sheet.sheet_properties.tabColor = GRONN
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"Skrev {OUT.relative_to(ROOT)} · {len(tomter)} tomter · {rb - 5} bilde-/videorader")
